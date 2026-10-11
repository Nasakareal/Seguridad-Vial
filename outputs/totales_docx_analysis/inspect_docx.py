import hashlib
import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document


def iter_table_text(table):
    for row_index, row in enumerate(table.rows):
        yield {
            "row": row_index,
            "cells": [cell.text for cell in row.cells],
        }


def main(path_text):
    path = Path(path_text)
    document = Document(path)
    result = {
        "path": str(path),
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "size": path.stat().st_size,
        "sections": [],
        "paragraphs": [
            {"index": index, "style": paragraph.style.name, "text": paragraph.text}
            for index, paragraph in enumerate(document.paragraphs)
            if paragraph.text.strip()
        ],
        "tables": [
            {
                "index": index,
                "rows": list(iter_table_text(table)),
            }
            for index, table in enumerate(document.tables)
        ],
        "headers": [],
        "footers": [],
        "markers": {},
        "parts": [],
    }

    for section_index, section in enumerate(document.sections):
        result["sections"].append(
            {
                "index": section_index,
                "width": section.page_width,
                "height": section.page_height,
                "top_margin": section.top_margin,
                "right_margin": section.right_margin,
                "bottom_margin": section.bottom_margin,
                "left_margin": section.left_margin,
                "header_distance": section.header_distance,
                "footer_distance": section.footer_distance,
                "orientation": str(section.orientation),
            }
        )
        result["headers"].append(
            [paragraph.text for paragraph in section.header.paragraphs if paragraph.text.strip()]
        )
        result["footers"].append(
            [paragraph.text for paragraph in section.footer.paragraphs if paragraph.text.strip()]
        )

    marker_pattern = re.compile(r"\$[A-Za-z][A-Za-z0-9_]*")
    with zipfile.ZipFile(path) as package:
        result["parts"] = sorted(package.namelist())
        for name in package.namelist():
            if not name.endswith((".xml", ".rels")):
                continue
            text = package.read(name).decode("utf-8", errors="ignore")
            markers = sorted(set(marker_pattern.findall(text)))
            if markers:
                result["markers"][name] = markers

    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main(sys.argv[1])
